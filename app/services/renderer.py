"""Document rendering — ADR-7 decision executed (REFERENCE_GAP_ANALYSIS A12/A2/A5).

ADR-7 (2026-09-30): reportlab (BSD) for PDF writing + python-docx (MIT) for
DOCX writing. Chosen over weasyprint (heavy system deps) and PyMuPDF (AGPL —
forbidden by R-TRUTH-6). Local-only rendering; no network; nothing auto-sends.

Templates are deliberately simple and ATS-safe: single column, standard fonts,
no tables-for-layout, no headers/footers that parsers choke on.
"""

from __future__ import annotations

import io
from typing import Any

__all__ = [
    "render_resume_pdf",
    "render_resume_docx",
    "render_prep_pack_docx",
    "RenderError",
]


class RenderError(ValueError):
    """Raised when a document cannot be rendered (missing section data)."""


# --------------------------------------------------------------------- PDF --
def render_resume_pdf(doc: dict[str, Any]) -> bytes:
    """Render the tailored resume as PDF (single column, ATS-safe).

    Expected ``doc`` shape (all strings come from confirmed facts / KB entries):
        contact: {name, email, phone, location}
        summary: str
        skills: [str]
        experience: [{role, company, period, bullets: [str]}]
        projects: [{title, bullets: [str]}]        (optional)
        education: [{degree, school, period}]      (optional)
        certifications: [str]                      (optional)
    """
    from reportlab.lib.colors import Color
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
    )

    contact = doc.get("contact") or {}
    name = str(contact.get("name") or "").strip()
    if not name:
        raise RenderError("resume needs a contact name (confirmed profile fact)")

    styles = {
        "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=16, leading=19),
        "contact": ParagraphStyle(
            "contact",
            fontName="Helvetica",
            fontSize=9.5,
            leading=12,
            textColor=Color(0.27, 0.27, 0.27),
        ),
        "section": ParagraphStyle(
            "section", fontName="Helvetica-Bold", fontSize=11.5, leading=14, spaceBefore=10
        ),
        "body": ParagraphStyle(
            "body",
            fontName="Helvetica",
            fontSize=10,
            leading=13,
            alignment=TA_LEFT,
            spaceBefore=2,
        ),
        "bullet": ParagraphStyle(
            "bullet", fontName="Helvetica", fontSize=10, leading=13, leftIndent=12, bulletIndent=4
        ),
        "role": ParagraphStyle(
            "role", fontName="Helvetica-Bold", fontSize=10.5, leading=13, spaceBefore=6
        ),
    }

    def esc(text: Any) -> str:
        return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"Resume - {name}",
        author=name,
    )
    story: list[Any] = [Paragraph(esc(name), styles["name"])]

    contact_bits = [
        str(contact.get(k) or "").strip()
        for k in ("email", "phone", "location")
        if str(contact.get(k) or "").strip()
    ]
    if contact_bits:
        story.append(Paragraph(esc(" · ".join(contact_bits)), styles["contact"]))
    if doc.get("summary"):
        story.append(Spacer(1, 4))
        story.append(Paragraph(esc(doc["summary"]), styles["body"]))

    if doc.get("skills"):
        story.append(Paragraph("Skills", styles["section"]))
        story.append(Paragraph(esc(", ".join(doc["skills"])), styles["body"]))

    for exp in doc.get("experience") or []:
        story.append(
            Paragraph(
                f"{esc(exp.get('role', ''))} — {esc(exp.get('company', ''))}"
                + (f" ({esc(exp.get('period', ''))})" if exp.get("period") else ""),
                styles["role"],
            )
        )
        for b in exp.get("bullets") or []:
            story.append(Paragraph(f"• {esc(b)}", styles["bullet"]))

    for proj in doc.get("projects") or []:
        story.append(Paragraph(esc(proj.get("title", "")), styles["role"]))
        for b in proj.get("bullets") or []:
            story.append(Paragraph(f"• {esc(b)}", styles["bullet"]))

    if doc.get("education"):
        story.append(Paragraph("Education", styles["section"]))
        for ed in doc["education"]:
            line = f"{ed.get('degree', '')} — {ed.get('school', '')}"
            if ed.get("period"):
                line += f" ({ed['period']})"
            story.append(Paragraph(esc(line), styles["body"]))

    if doc.get("certifications"):
        story.append(Paragraph("Certifications", styles["section"]))
        for c in doc["certifications"]:
            story.append(Paragraph(f"• {esc(c)}", styles["bullet"]))

    pdf.build(story)
    return buf.getvalue()


# -------------------------------------------------------------------- DOCX --
def render_resume_docx(doc: dict[str, Any]) -> bytes:
    """Same resume structure as render_resume_pdf, emitted as .docx (python-docx)."""
    import docx
    from docx.shared import Pt

    contact = doc.get("contact") or {}
    name = str(contact.get("name") or "").strip()
    if not name:
        raise RenderError("resume needs a contact name (confirmed profile fact)")

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    h = d.add_paragraph()
    run = h.add_run(name)
    run.bold = True
    run.font.size = Pt(16)

    contact_bits = [
        str(contact.get(k) or "").strip()
        for k in ("email", "phone", "location")
        if str(contact.get(k) or "").strip()
    ]
    if contact_bits:
        c = d.add_paragraph(contact_bits[-1] if False else " · ".join(contact_bits))
        for r in c.runs:
            r.font.size = Pt(9.5)

    if doc.get("summary"):
        d.add_paragraph(str(doc["summary"]))

    def section(title: str) -> None:
        p = d.add_paragraph()
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(11.5)

    if doc.get("skills"):
        section("Skills")
        d.add_paragraph(", ".join(str(s) for s in doc["skills"]))

    for exp in doc.get("experience") or []:
        p = d.add_paragraph()
        r = p.add_run(f"{exp.get('role', '')} — {exp.get('company', '')}")
        r.bold = True
        if exp.get("period"):
            p.add_run(f" ({exp['period']})")
        for b in exp.get("bullets") or []:
            d.add_paragraph(str(b), style="List Bullet")

    for proj in doc.get("projects") or []:
        p = d.add_paragraph()
        p.add_run(str(proj.get("title", ""))).bold = True
        for b in proj.get("bullets") or []:
            d.add_paragraph(str(b), style="List Bullet")

    if doc.get("education"):
        section("Education")
        for ed in doc["education"]:
            line = f"{ed.get('degree', '')} — {ed.get('school', '')}"
            if ed.get("period"):
                line += f" ({ed['period']})"
            d.add_paragraph(line)

    if doc.get("certifications"):
        section("Certifications")
        for c in doc["certifications"]:
            d.add_paragraph(str(c), style="List Bullet")

    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


# --------------------------------------------------------------- prep pack --
def render_prep_pack_docx(pack: dict[str, Any]) -> bytes:
    """Interview-prep pack (A2): JD-derived questions + fact-grounded answer outlines.

    Expected shape:
        job_title, company, prepared_at
        questions: [{question, outline, source}]
        talking_points: [str]
        questions_to_ask: [str]
    """
    import docx
    from docx.shared import Pt

    job_title = str(pack.get("job_title") or "").strip()
    company = str(pack.get("company") or "").strip()
    if not job_title or not company:
        raise RenderError("prep pack needs job_title and company")

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    h = d.add_paragraph()
    r = h.add_run(f"Interview Prep — {job_title} @ {company}")
    r.bold = True
    r.font.size = Pt(15)
    if pack.get("prepared_at"):
        meta = d.add_paragraph(str(pack["prepared_at"]))
        for run in meta.runs:
            run.font.size = Pt(9)

    def section(title: str) -> None:
        p = d.add_paragraph()
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(11.5)

    questions = pack.get("questions") or []
    if questions:
        section("Likely questions & answer outlines")
        for i, q in enumerate(questions, start=1):
            p = d.add_paragraph()
            run = p.add_run(f"{i}. {q.get('question', '')}")
            run.bold = True
            if q.get("outline"):
                ol = d.add_paragraph(q["outline"])
                ol.paragraph_format.left_indent = Pt(18)
            if q.get("source"):
                s = d.add_paragraph(f"grounded in: {q['source']}")
                s.paragraph_format.left_indent = Pt(18)
                for run in s.runs:
                    run.italic = True
                    run.font.size = Pt(9)

    talking = pack.get("talking_points") or []
    if talking:
        section("Your talking points")
        for t in talking:
            d.add_paragraph(str(t), style="List Bullet")

    ask = pack.get("questions_to_ask") or []
    if ask:
        section("Questions to ask them")
        for a in ask:
            d.add_paragraph(str(a), style="List Bullet")

    out = io.BytesIO()
    d.save(out)
    return out.getvalue()
